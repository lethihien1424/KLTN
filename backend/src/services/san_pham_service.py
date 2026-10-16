### D:\KLTN\KLTN\backend\src\services\san_pham_service.py
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from pydantic import ValidationError
from sqlalchemy import (
    func,
    or_,
    select,
)
from sqlalchemy.exc import (
    IntegrityError,
    SQLAlchemyError,
)

from src.models.san_pham_model import SanPham
from src.schemas.san_pham_schema import (
    SanPhamListResponse,
    SanPhamRequest,
    SanPhamResponse,
)
from src.services.data_processing.kiem_tra_du_lieu import (
    to_decimal,
    to_string,
)


HEADERS = [
    "ten_san_pham",
    "khoi_luong",
    "don_vi_do_luong",
    "nhom_san_pham",
    "trang_thai",
    "don_gia",
    "don_vi_tien_te",
]

DEFAULT_GROUPS = [
    "CA_PHE_BOT_TRUYEN_THONG",
    "CA_PHE_BOT_HUONG_VI",
    "CA_PHE_HAT",
    "CA_PHE_BOT",
]

VALIDATION_REASONS = {
    "missing": "bắt buộc nhập",
    "string_too_short": "không được để trống",
    "string_too_long": "vượt quá độ dài cho phép",
    "greater_than": "phải lớn hơn 0",
    "greater_than_equal": "không được âm",
    "decimal_max_digits": "vượt quá số chữ số cho phép",
    "decimal_max_places": (
        "quá số chữ số thập phân cho phép"
    ),
    "literal_error": (
        "giá trị không hợp lệ; "
        "kiểm tra hướng dẫn template"
    ),
}


class SanPhamService:
    @staticmethod
    def statement(
        search=None,
        nhom=None,
        trang_thai=None,
    ):
        statement = select(SanPham)

        if search and search.strip():
            keyword = search.strip().lower()

            statement = statement.where(
                or_(
                    func.lower(
                        SanPham.ma_san_pham
                    ).contains(
                        keyword,
                        autoescape=True,
                    ),
                    func.lower(
                        SanPham.ten_san_pham
                    ).contains(
                        keyword,
                        autoescape=True,
                    ),
                )
            )

        if nhom:
            statement = statement.where(
                SanPham.nhom_san_pham == nhom
            )

        if trang_thai:
            statement = statement.where(
                SanPham.trang_thai == trang_thai
            )

        return statement

    @classmethod
    def get_list(
        cls,
        db,
        search,
        nhom,
        trang_thai,
        page,
        page_size,
    ) -> SanPhamListResponse:
        statement = cls.statement(
            search,
            nhom,
            trang_thai,
        )

        total = db.scalar(
            select(func.count()).select_from(
                statement.subquery()
            )
        ) or 0

        rows = db.scalars(
            statement
            .order_by(SanPham.ma_san_pham)
            .offset((page - 1) * page_size)
            .limit(page_size)
        ).all()

        return SanPhamListResponse(
            items=[
                SanPhamResponse.model_validate(row)
                for row in rows
            ],
            total=total,
            page=page,
            page_size=page_size,
        )

    @staticmethod
    def get_one(
        db,
        ma,
        lock=False,
    ):
        statement = select(SanPham).where(
            SanPham.ma_san_pham == ma
        )

        if lock:
            statement = statement.with_for_update()

        row = db.scalar(statement)

        if row is None:
            raise HTTPException(
                status_code=404,
                detail="Không tìm thấy sản phẩm.",
            )

        return row

    @staticmethod
    def groups(db):
        groups = db.scalars(
            select(SanPham.nhom_san_pham)
            .where(
                SanPham.nhom_san_pham.is_not(None)
            )
            .distinct()
        ).all()

        return sorted(
            set(DEFAULT_GROUPS).union(groups)
        )

    @staticmethod
    def ensure_name(
        db,
        name,
        exclude=None,
    ):
        statement = select(
            SanPham.ma_san_pham
        ).where(
            func.lower(SanPham.ten_san_pham)
            == name.lower()
        )

        if exclude:
            statement = statement.where(
                SanPham.ma_san_pham != exclude
            )

        if db.scalar(statement.limit(1)) is not None:
            raise HTTPException(
                status_code=409,
                detail=(
                    f'Sản phẩm "{name}" đã tồn tại. '
                    "Hãy chỉnh sửa sản phẩm hiện có."
                ),
            )

    @classmethod
    def save(
        cls,
        db,
        payload: SanPhamRequest,
        ma=None,
    ):
        try:
            cls.ensure_name(
                db,
                payload.ten_san_pham,
                ma,
            )

            if ma:
                row = cls.get_one(
                    db,
                    ma,
                    lock=True,
                )

                for field, value in (
                    payload.model_dump().items()
                ):
                    setattr(row, field, value)

                row.thoi_gian_cap_nhat = (
                    datetime.now(timezone.utc)
                )
            else:
                # Không truyền mã.
                # DEFAULT của PostgreSQL tự sinh mã.
                row = SanPham(
                    **payload.model_dump()
                )

                db.add(row)

            db.flush()
            db.refresh(row)

            result = SanPhamResponse.model_validate(
                row
            )

            db.commit()

            return result

        except HTTPException:
            db.rollback()
            raise

        except IntegrityError as exc:
            db.rollback()

            raise HTTPException(
                status_code=409,
                detail=(
                    "Không thể lưu sản phẩm vì dữ liệu "
                    "trùng hoặc vi phạm ràng buộc DB."
                ),
            ) from exc

        except SQLAlchemyError as exc:
            db.rollback()

            raise HTTPException(
                status_code=500,
                detail=(
                    "Không thể lưu sản phẩm lúc này."
                ),
            ) from exc

        except Exception:
            db.rollback()
            raise

    @classmethod
    def delete(
        cls,
        db,
        ma,
    ):
        try:
            row = cls.get_one(
                db,
                ma,
                lock=True,
            )

            db.delete(row)
            db.commit()

        except HTTPException:
            db.rollback()
            raise

        except IntegrityError as exc:
            db.rollback()

            raise HTTPException(
                status_code=409,
                detail=(
                    "Sản phẩm đang được dùng trong "
                    "dữ liệu khác. Hãy chuyển sang "
                    "Ngưng kinh doanh thay vì xóa."
                ),
            ) from exc

        except SQLAlchemyError as exc:
            db.rollback()

            raise HTTPException(
                status_code=500,
                detail=(
                    "Không thể xóa sản phẩm lúc này."
                ),
            ) from exc

    @classmethod
    def import_rows(
        cls,
        db,
        frame,
    ):
        actual = set(frame.columns) - {
            "_row_number"
        }

        required = set(HEADERS) - {
            "nhom_san_pham"
        }

        missing = required - actual
        extra = actual - set(HEADERS)

        if missing or extra:
            raise HTTPException(
                status_code=422,
                detail={
                    "message": (
                        "File không đúng template sản phẩm."
                    ),
                    "thieu_cot": sorted(missing),
                    "du_cot": sorted(extra),
                },
            )

        models = []
        seen = set()

        try:
            for _, source in frame.iterrows():
                number = int(
                    source["_row_number"]
                )

                unit = to_string(
                    source["don_vi_do_luong"],
                    "don_vi_do_luong",
                    number,
                ).lower()

                weight = to_decimal(
                    source["khoi_luong"],
                    "khoi_luong",
                    number,
                )

                if unit == "g":
                    weight /= Decimal("1000")
                elif unit != "kg":
                    raise ValueError(
                        f"Dòng {number}, "
                        "cột don_vi_do_luong: "
                        "chỉ chấp nhận kg hoặc g."
                    )

                data = {
                    "ten_san_pham": to_string(
                        source["ten_san_pham"],
                        "ten_san_pham",
                        number,
                    ),
                    "khoi_luong": weight,
                    "don_vi_do_luong": "kg",
                    "nhom_san_pham": to_string(
                        source.get("nhom_san_pham"),
                        "nhom_san_pham",
                        number,
                        required=False,
                    ),
                    "trang_thai": to_string(
                        source["trang_thai"],
                        "trang_thai",
                        number,
                    ).upper(),
                    "don_gia": to_decimal(
                        source["don_gia"],
                        "don_gia",
                        number,
                    ),
                    "don_vi_tien_te": to_string(
                        source["don_vi_tien_te"],
                        "don_vi_tien_te",
                        number,
                    ).upper(),
                }

                try:
                    payload = (
                        SanPhamRequest.model_validate(
                            data
                        )
                    )
                except ValidationError as exc:
                    errors = "; ".join(
                        (
                            "cột "
                            + ".".join(
                                str(part)
                                for part in error["loc"]
                            )
                            + ": "
                            + VALIDATION_REASONS.get(
                                error["type"],
                                error["msg"],
                            )
                        )
                        for error in exc.errors()
                    )

                    raise ValueError(
                        f"Dòng {number}: {errors}"
                    ) from exc

                key = (
                    payload.ten_san_pham.casefold()
                )

                if key in seen:
                    raise ValueError(
                        f"Dòng {number}, "
                        "cột ten_san_pham: "
                        "tên sản phẩm bị lặp trong file."
                    )

                seen.add(key)

                try:
                    cls.ensure_name(
                        db,
                        payload.ten_san_pham,
                    )
                except HTTPException as exc:
                    raise ValueError(
                        f"Dòng {number}, "
                        "cột ten_san_pham: "
                        f"{exc.detail}"
                    ) from exc

                models.append(
                    SanPham(
                        **payload.model_dump()
                    )
                )

            # Chỉ thêm khi tất cả dòng đều hợp lệ.
            db.add_all(models)
            db.flush()
            db.commit()

            return {
                "success": True,
                "message": (
                    f"Đã nhập {len(models)} sản phẩm."
                ),
                "total_rows": len(frame),
                "imported_rows": len(models),
            }

        except ValueError as exc:
            db.rollback()

            raise HTTPException(
                status_code=422,
                detail=str(exc),
            ) from exc

        except HTTPException:
            db.rollback()
            raise

        except IntegrityError as exc:
            db.rollback()

            raise HTTPException(
                status_code=409,
                detail=(
                    "Không nhập dữ liệu: file vi phạm "
                    "ràng buộc DB."
                ),
            ) from exc

        except SQLAlchemyError as exc:
            db.rollback()

            raise HTTPException(
                status_code=500,
                detail=(
                    "Không thể nhập sản phẩm lúc này."
                ),
            ) from exc

        except Exception:
            db.rollback()
            raise

    @staticmethod
    def excel(rows):
        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "SanPham"
        sheet.append(HEADERS)

        for row in rows:
            sheet.append([
                row.ten_san_pham,
                float(row.khoi_luong),
                row.don_vi_do_luong,
                row.nhom_san_pham or "",
                row.trang_thai,
                float(row.don_gia),
                row.don_vi_tien_te,
            ])

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 32

        for cell in sheet[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                "solid",
                fgColor="2563EB",
            )
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="center",
            )

        widths = {
            "A": 48,
            "B": 16,
            "C": 22,
            "D": 32,
            "E": 26,
            "F": 20,
            "G": 20,
        }

        for column, width in widths.items():
            sheet.column_dimensions[
                column
            ].width = width

        for row in sheet.iter_rows(min_row=2):
            row[1].number_format = "0.000"
            row[5].number_format = "#,##0.00"

        if rows:
            report = workbook.create_sheet(
                "DanhSach"
            )

            report.append([
                "Mã sản phẩm",
                "Tên sản phẩm",
                "Khối lượng (g)",
                "Nhóm",
                "Trạng thái",
                "Đơn giá (VND)",
                "Thời gian tạo",
                "Thời gian cập nhật",
            ])

            for row in rows:
                grams = row.khoi_luong * (
                    1000
                    if row.don_vi_do_luong == "kg"
                    else 1
                )

                report.append([
                    row.ma_san_pham,
                    row.ten_san_pham,
                    float(grams),
                    row.nhom_san_pham or "",
                    row.trang_thai,
                    float(row.don_gia),
                    row.thoi_gian_tao.isoformat(),
                    row.thoi_gian_cap_nhat.isoformat(),
                ])

            report.freeze_panes = "A2"

            for column in report.columns:
                report.column_dimensions[
                    column[0].column_letter
                ].width = 28

            report.column_dimensions[
                "B"
            ].width = 48

        guide = workbook.create_sheet(
            "HuongDan"
        )

        instructions = [
            (
                "Nhập dữ liệu ở sheet SanPham. "
                "Không đổi tên cột."
            ),
            (
                "Không nhập mã sản phẩm hoặc thời gian; "
                "hệ thống tự sinh."
            ),
            (
                "Khối lượng: 0.250 và kg, hoặc 250 và g. "
                "DB lưu kg."
            ),
            (
                "Trạng thái: DANG_KINH_DOANH "
                "hoặc NGUNG_BAN."
            ),
            (
                "Đơn giá là số, ví dụ 125000. "
                "Tiền tệ: VND."
            ),
            "Nhóm sản phẩm có thể để trống.",
            (
                "Import dùng để thêm sản phẩm mới. "
                "Sản phẩm đã tồn tại phải dùng "
                "chức năng Sửa."
            ),
        ]

        for line in instructions:
            guide.append([line])

        guide.column_dimensions[
            "A"
        ].width = 110

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        return buffer